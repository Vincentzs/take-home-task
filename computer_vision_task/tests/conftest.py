"""Shared fixtures for pipeline tests."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document

MakeSopDocx = Callable[..., Path]


@pytest.fixture
def make_sop_docx(tmp_path: Path) -> MakeSopDocx:
    def _make(rows: list[tuple[str, ...]], header_label: str = "Design Limits") -> Path:
        doc = Document()
        doc.add_heading("Operating Limits", level=1)
        doc.add_paragraph(header_label + ":")
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=1 + len(rows), cols=cols)
        # First row: column headers (Pressure (psig), Temperature (°F)) in cols 1+
        if cols >= 3:
            table.rows[0].cells[1].text = "Pressure (psig)"
            table.rows[0].cells[2].text = "Temperature (°F)"
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                table.rows[i].cells[j].text = value
        out = tmp_path / "sop.docx"
        doc.save(str(out))
        return out

    return _make
